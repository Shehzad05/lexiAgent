# """
# SOLID:
#   S — one job: convert any document to text
# """
# from __future__ import annotations
# import io, os, tempfile

# def load_document(data: bytes, filename: str) -> str:
#     """Convert any supported document to plain text / markdown."""
#     ext = os.path.splitext(filename)[1].lower()

#     # ── Docling: handles PDF, DOCX, PPTX, XLSX, HTML, images ────────────────
#     try:
#         from docling.document_converter import DocumentConverter
#         with tempfile.NamedTemporaryFile(suffix=ext or ".pdf", delete=False) as tmp:
#             tmp.write(data)
#             tmp_path = tmp.name
#         result = DocumentConverter().convert(tmp_path)
#         os.unlink(tmp_path)
#         return result.document.export_to_markdown()
#     except ImportError:
#         pass
#     except Exception as e:
#         print(f"[DocumentLoader] Docling error: {e} — using fallback")

#     # ── PDF fallback ─────────────────────────────────────────────────────────
#     if ext == ".pdf":
#         try:
#             import pdfplumber
#             with pdfplumber.open(io.BytesIO(data)) as pdf:
#                 return "\n".join(p.extract_text() or "" for p in pdf.pages)
#         except ImportError:
#             pass
#         try:
#             import PyPDF2
#             r = PyPDF2.PdfReader(io.BytesIO(data))
#             return "\n".join(p.extract_text() or "" for p in r.pages)
#         except ImportError:
#             pass

#     # ── DOCX fallback ────────────────────────────────────────────────────────
#     if ext == ".docx":
#         try:
#             from docx import Document
#             return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
#         except ImportError:
#             pass

#     # ── Plain text fallback ──────────────────────────────────────────────────
#     return data.decode("utf-8", errors="ignore")
import io
import os

def load_document(data: bytes, filename: str) -> str:
    """Convert any supported document to plain text."""
    ext = os.path.splitext(filename)[1].lower()

    # PDF
    if ext == ".pdf":
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                return "\n".join(p.extract_text() or "" for p in pdf.pages).strip()
        except Exception:
            pass
        try:
            import PyPDF2
            r = PyPDF2.PdfReader(io.BytesIO(data))
            return "\n".join(p.extract_text() or "" for p in r.pages).strip()
        except Exception:
            pass

    # DOCX
    if ext == ".docx":
        try:
            from docx import Document
            return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs).strip()
        except Exception:
            pass

    # Plain text fallback
    return data.decode("utf-8", errors="ignore")